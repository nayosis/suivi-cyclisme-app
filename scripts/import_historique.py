"""Import unique des fichiers Excel historiques (2016-2026) vers la base SQLite.

Structure des fichiers source (suivi perf velo/*.xlsx) : un onglet par mois,
avec des colonnes qui varient selon les annees :
  - 2016-2018 : pas de colonne Denivele+
  - 2019-2020 : Denivele+ present, mais un seul velo (pas de colonne velo)
  - 2021-2026 : Denivele+ present, deux colonnes velo (Brochard, S1NEO) —
    le velo utilise est celui dont la colonne contient la distance de la sortie.

Regles appliquees (validees avec l'utilisateur) :
  - Si aucune colonne velo n'existe (2016-2020) -> velo "Brochard".
  - Si les colonnes velo existent mais qu'aucune n'est renseignee pour une
    ligne -> on prend la premiere colonne velo (Brochard).
  - Les lignes sans date, sans distance numerique, ou avec distance <= 0
    (repos, cases vides) sont ignorees.
  - Le script est idempotent : une ligne deja presente (meme date/nom/
    distance/velo) n'est pas reimportee, donc on peut le relancer sans crainte.
"""

import datetime
import glob
import sqlite3
import unicodedata

import openpyxl

from app.db import get_connection, init_db

FALLBACK_BIKE_NAME = "Brochard"
SOURCE_GLOB = r"c:\dev\src\suivi-cyclisme-app\suivi perf velo\*.xlsx"


def norm(s) -> str:
    if s is None:
        return ""
    text = unicodedata.normalize("NFKD", str(s))
    text = "".join(c for c in text if not unicodedata.combining(c))
    return text.lower().strip()


def find_col(header: tuple, needle: str) -> int | None:
    for idx, val in enumerate(header):
        if needle in norm(val):
            return idx
    return None


def detect_columns(header: tuple) -> dict | None:
    date_col = find_col(header, "date")
    name_col = find_col(header, "designation")
    distance_col = find_col(header, "distance")
    duree_col = find_col(header, "duree")
    remarque_col = find_col(header, "remarque")
    if None in (date_col, name_col, distance_col, duree_col, remarque_col):
        return None
    elevation_col = find_col(header, "dev")
    bike_cols = [
        (idx, str(val).strip())
        for idx, val in enumerate(header)
        if idx > remarque_col and val is not None and str(val).strip()
    ]
    return {
        "date": date_col,
        "name": name_col,
        "distance": distance_col,
        "duration_total": duree_col + 2,
        "duration_hours": duree_col,
        "duration_minutes": duree_col + 1,
        "elevation": elevation_col,
        "remarque": remarque_col,
        "bike_cols": bike_cols,
    }


def get_or_create_bike(conn: sqlite3.Connection, cache: dict, name: str) -> int:
    if name in cache:
        return cache[name]
    row = conn.execute("SELECT id FROM bikes WHERE name = ?", (name,)).fetchone()
    if row:
        bike_id = row["id"]
    else:
        cur = conn.execute("INSERT INTO bikes (name) VALUES (?)", (name,))
        bike_id = cur.lastrowid
    cache[name] = bike_id
    return bike_id


def pick_bike(cols: dict, row: tuple, distance: float) -> tuple[str, bool]:
    """Retourne (nom_du_velo, utilisation_du_repli)."""
    bike_cols = cols["bike_cols"]
    if not bike_cols:
        return FALLBACK_BIKE_NAME, True

    candidates = []
    for idx, name in bike_cols:
        val = row[idx] if idx < len(row) else None
        if isinstance(val, (int, float)) and val:
            candidates.append((name, val))

    if len(candidates) == 1:
        return candidates[0][0], False
    if len(candidates) > 1:
        matching = [n for n, v in candidates if abs(v - distance) < 0.01]
        if matching:
            return matching[0], False
        return candidates[0][0], False
    # aucune colonne velo renseignee -> repli sur la premiere colonne
    return bike_cols[0][1], True


def import_file(conn: sqlite3.Connection, path: str, bike_cache: dict, stats: dict) -> None:
    wb = openpyxl.load_workbook(path, data_only=True)
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        if len(rows) < 2:
            continue
        header = rows[1]
        cols = detect_columns(header)
        if cols is None:
            continue  # onglet non standard (Feuil1/Feuil2 vides, etc.)

        for row in rows[2:]:
            stats["scanned"] += 1

            date_val = row[cols["date"]] if cols["date"] < len(row) else None
            if not isinstance(date_val, (datetime.datetime, datetime.date)):
                stats["skipped_no_date"] += 1
                continue

            distance_val = row[cols["distance"]] if cols["distance"] < len(row) else None
            if not isinstance(distance_val, (int, float)) or distance_val <= 0:
                stats["skipped_no_distance"] += 1
                continue

            name_val = row[cols["name"]] if cols["name"] < len(row) else None
            name_val = str(name_val).strip() if name_val else "Sortie"

            total_idx = cols["duration_total"]
            hours_idx = cols["duration_hours"]
            minutes_idx = cols["duration_minutes"]
            total_val = row[total_idx] if total_idx < len(row) else None
            if isinstance(total_val, (int, float)):
                duration_minutes = int(round(total_val))
            else:
                h = row[hours_idx] if hours_idx < len(row) and isinstance(row[hours_idx], (int, float)) else 0
                m = row[minutes_idx] if minutes_idx < len(row) and isinstance(row[minutes_idx], (int, float)) else 0
                duration_minutes = int(round(h * 60 + m))

            elevation_col = cols["elevation"]
            elevation_val = 0
            if elevation_col is not None and elevation_col < len(row):
                raw = row[elevation_col]
                if isinstance(raw, (int, float)):
                    elevation_val = int(round(raw))

            remarque_idx = cols["remarque"]
            comment_val = row[remarque_idx] if remarque_idx < len(row) else None
            comment_val = str(comment_val).strip() if comment_val else None

            bike_name, used_fallback = pick_bike(cols, row, distance_val)
            if used_fallback:
                stats["fallback_bike"] += 1
            bike_id = get_or_create_bike(conn, bike_cache, bike_name)

            date_iso = date_val.strftime("%Y-%m-%d")

            existing = conn.execute(
                """
                SELECT id FROM rides
                WHERE date = ? AND name = ? AND distance_km = ? AND bike_id = ?
                """,
                (date_iso, name_val, distance_val, bike_id),
            ).fetchone()
            if existing:
                stats["skipped_duplicate"] += 1
                continue

            conn.execute(
                """
                INSERT INTO rides (date, name, duration_minutes, distance_km, elevation_m, comment, bike_id)
                VALUES (?, ?, ?, ?, ?, ?, ?)
                """,
                (date_iso, name_val, duration_minutes, distance_val, elevation_val, comment_val, bike_id),
            )
            stats["imported"] += 1


def main() -> None:
    conn = get_connection()
    init_db(conn)

    bike_cache: dict[str, int] = {}
    stats = {
        "scanned": 0,
        "imported": 0,
        "skipped_no_date": 0,
        "skipped_no_distance": 0,
        "skipped_duplicate": 0,
        "fallback_bike": 0,
    }

    files = sorted(glob.glob(SOURCE_GLOB))
    for path in files:
        import_file(conn, path, bike_cache, stats)
        conn.commit()

    print(f"Fichiers traites : {len(files)}")
    print(f"Lignes scannees  : {stats['scanned']}")
    print(f"Sorties importees: {stats['imported']}")
    print(f"Doublons ignores : {stats['skipped_duplicate']}")
    print(f"Ignorees (pas de date)    : {stats['skipped_no_date']}")
    print(f"Ignorees (pas de distance): {stats['skipped_no_distance']}")
    print(f"Repli sur 1er velo (Brochard) : {stats['fallback_bike']}")
    print("Velos:", bike_cache)


if __name__ == "__main__":
    main()
